import tkinter as tk
from tkinter import filedialog, messagebox
import os
from pdf2image import convert_from_path
import pytesseract
from docx import Document
from docx.shared import Inches
from PIL import Image

# GUI setup
root = tk.Tk()
root.title("PDF to Word Converter with OCR")
root.geometry("500x300")

# Variables
pdf_path_var = tk.StringVar()
output_dir_var = tk.StringVar()
ocr_var = tk.BooleanVar()

# Function to select PDF file
def select_pdf():
    file_path = filedialog.askopenfilename(
        filetypes=[("PDF Files", "*.pdf")],
        title="Select PDF File"
    )
    if file_path:
        pdf_path_var.set(file_path)

# Function to select output directory
def select_output_dir():
    dir_path = filedialog.askdirectory(title="Select Output Directory")
    if dir_path:
        output_dir_var.set(dir_path)

# Function to convert PDF
def convert_pdf():
    pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    
    pdf_path = pdf_path_var.get()
    output_dir = output_dir_var.get()
    use_ocr = ocr_var.get()

    if not pdf_path or not os.path.isfile(pdf_path):
        messagebox.showerror("Error", "Please select a valid PDF file.")
        return
    if not output_dir or not os.path.isdir(output_dir):
        messagebox.showerror("Error", "Please select a valid output directory.")
        return

    try:
        output_word_path = os.path.join(output_dir, "Converted_Output.docx")
        document = Document()

        # Convert each page to image
        images = convert_from_path(pdf_path, dpi=300)

        for i, img in enumerate(images):
            temp_img_path = os.path.join(output_dir, f"page_{i}.png")
            img.save(temp_img_path)

            section = document.sections[0]
            available_width = section.page_width - section.left_margin - section.right_margin
            document.add_picture(temp_img_path, width=available_width)

            if use_ocr:
                text = pytesseract.image_to_string(img, lang="eng")
                document.add_paragraph(text)

            if i < len(images) - 1:
                document.add_page_break()

        document.save(output_word_path)
        messagebox.showinfo("Success", f"File saved to: {output_word_path}")

    except Exception as e:
        messagebox.showerror("Error", f"An error occurred: {str(e)}")

# Layout
tk.Label(root, text="Select PDF File:").pack(anchor="w", padx=10, pady=5)
tk.Entry(root, textvariable=pdf_path_var, width=50).pack(anchor="w", padx=10)
tk.Button(root, text="Browse", command=select_pdf).pack(anchor="w", padx=10)

tk.Label(root, text="Select Output Directory:").pack(anchor="w", padx=10, pady=5)
tk.Entry(root, textvariable=output_dir_var, width=50).pack(anchor="w", padx=10)
tk.Button(root, text="Browse", command=select_output_dir).pack(anchor="w", padx=10)

tk.Checkbutton(root, text="Use OCR (for scanned PDFs)", variable=ocr_var).pack(anchor="w", padx=10, pady=5)

tk.Button(root, text="Convert", command=convert_pdf, bg="blue", fg="white").pack(pady=20)

# Show GUI
root.mainloop()

